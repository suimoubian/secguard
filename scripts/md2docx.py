# -*- coding: utf-8 -*-
"""Markdown -> DOCX 转换器(为项目说明文档定制)

为什么走这条路:直接在 Python 源码里写中文长文本,全角引号一旦被编辑器
转成 ASCII 引号就会截断字符串。把内容放在 Markdown 里,引号不再有语法意义,
这个问题就彻底消失了。

支持的语法:
  # / ## / ###     标题
  | a | b |        表格(第二行是 |---|---| 分隔行)
  ```              代码块
  - 项目           无序列表
  1. 项目          有序列表
  > 标题 / > 正文  强调框
  ---              分页
  **粗体**         行内加粗
"""

import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN = "微软雅黑"
MONO = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
ACCENT2 = RGBColor(0x2E, 0x74, 0xB5)
GREY = RGBColor(0x59, 0x59, 0x59)
DARK = RGBColor(0x22, 0x22, 0x22)


# --------------------------------------------------------------------------
def set_font(run, name=CN, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def add_bottom_border(p, color="1F4E79", sz="6"):
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_CODE = re.compile(r"`(.+?)`")
INLINE_LINK = re.compile(r"\[([^\]]+)\]\([^)]*\)")


def add_rich_text(p, text, size=10.5, color=DARK, base_bold=False):
    """处理 **粗体**、`代码`、[文字](链接)"""
    # 链接只保留文字:文档里链接的 label 通常就是文件名本身,
    # 打印出来带一长串 URL 反而更难看。
    text = INLINE_LINK.sub(r"\1", text)

    tokens = []
    pos = 0
    pattern = re.compile(r"\*\*(.+?)\*\*|`(.+?)`")
    for m in pattern.finditer(text):
        if m.start() > pos:
            tokens.append((text[pos:m.start()], False, False))
        if m.group(1) is not None:
            tokens.append((m.group(1), True, False))
        else:
            tokens.append((m.group(2), False, True))
        pos = m.end()
    if pos < len(text):
        tokens.append((text[pos:], False, False))

    for txt, bold, code in tokens:
        if not txt:
            continue
        r = p.add_run(txt)
        if code:
            set_font(r, name=MONO, size=size - 1, bold=False,
                     color=RGBColor(0xB0, 0x30, 0x60))
        else:
            set_font(r, size=size, bold=(bold or base_bold), color=color)


def _mark_header_row(row):
    """让表头在跨页时重复出现"""
    trpr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trpr.append(hdr)


def _no_split(row):
    """禁止一行被拆到两页"""
    trpr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    cs.set(qn("w:val"), "true")
    trpr.append(cs)


def _set_fixed_layout(t):
    tblpr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)


def _display_width(s):
    """粗略估算显示宽度:中文算 2,ASCII 算 1"""
    return sum(2 if ord(ch) > 127 else 1 for ch in s)


def _set_grid(t, widths_cm):
    """固定布局下,Word 真正依据的是 tblGrid 里的 gridCol 宽度,
    只设 cell.width 是不够的(那样只有等分效果)。"""
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is None:
        return
    cols = grid.findall(qn("w:gridCol"))
    for col, w in zip(cols, widths_cm):
        col.set(qn("w:w"), str(int(w * 567)))   # 1cm = 567 twips


def compute_widths(rows, total_cm=16.0, min_cm=2.5):
    """按每列内容长度分配列宽,避免「可用」这种两字列占掉三分之一,
    同时给短列留足下限,不让「SQL 注入」这种词被挤到换行。"""
    ncol = len(rows[0])
    medians = []
    for c in range(ncol):
        vals = sorted(_display_width(r[c]) for r in rows if c < len(r))
        med = vals[len(vals) // 2] if vals else 4
        medians.append(max(med, 5))
    s = sum(medians)
    widths = [total_cm * m / s for m in medians]
    widths = [max(min_cm, w) for w in widths]
    scale = total_cm / sum(widths)
    return [w * scale for w in widths]


# --------------------------------------------------------------------------
class Builder:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.4)
        sec.right_margin = Cm(2.4)
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.2)

        st = self.doc.styles["Normal"]
        st.font.name = CN
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CN)

        # 页脚页码
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = "PAGE"
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        run._r.append(b); run._r.append(i); run._r.append(e)
        set_font(run, size=9, color=GREY)

    # ---- 封面 -------------------------------------------------------------
    def cover(self, lines):
        for _ in range(3):
            p = self.doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        for kind, text in lines:
            if kind == "h1":
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(text); set_font(r, size=44, bold=True, color=ACCENT)
            elif kind == "h2":
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(18)
                r = p.add_run(text); set_font(r, size=20, bold=True, color=DARK)
            elif kind == "h3":
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run(text); set_font(r, size=13, bold=True, color=ACCENT2)
            elif kind == "quote":
                self.callout(*text)
            else:  # p / text
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(text); set_font(r, size=11, color=GREY)

    # ---- 块元素 -----------------------------------------------------------
    def h1(self, text):
        p = self.doc.add_heading(level=1)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(text); set_font(r, size=17, bold=True, color=ACCENT)
        add_bottom_border(p)

    def h2(self, text):
        p = self.doc.add_heading(level=2)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text); set_font(r, size=13.5, bold=True, color=ACCENT)

    def h3(self, text):
        p = self.doc.add_heading(level=3)
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text); set_font(r, size=11.5, bold=True, color=ACCENT2)

    def para(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.5
        add_rich_text(p, text)

    def bullet(self, text):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.45
        p.paragraph_format.left_indent = Pt(22)
        add_rich_text(p, text)

    def num(self, text):
        p = self.doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.45
        p.paragraph_format.left_indent = Pt(22)
        add_rich_text(p, text)

    def callout(self, title, body):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.text = ""
        shade(cell, "EAF3FB")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(title); set_font(r, size=10.5, bold=True, color=ACCENT)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.45
        add_rich_text(p2, body, size=10)
        cell.width = Cm(16.2)
        # 表格后留一点空
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0)

    def code(self, lines):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.text = ""
        shade(cell, "F5F5F5")
        for idx, line in enumerate(lines):
            p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(line if line else " ")
            set_font(r, name=MONO, size=8.5, color=DARK)
        cell.width = Cm(16.2)
        # 短代码块(比如 ASCII 示意图)不允许跨页,否则图会被切成两半;
        # 长代码块必须允许跨页,否则超过一页会被截断。
        if len(lines) <= 32:
            _no_split(t.rows[0])
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0)

    def table(self, rows):
        header = rows[0]
        # 注意:分隔行(|---|)在 parse() 里已经剔除了,所以这里从第 1 行开始,
        # 不能写成 rows[2:] —— 那会把每张表的第一行数据也吃掉。
        body = rows[1:]
        ncol = len(header)
        t = self.doc.add_table(rows=1, cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _set_fixed_layout(t)

        widths = compute_widths(rows)

        hdr = t.rows[0].cells
        for i, txt in enumerate(header):
            hdr[i].text = ""
            p = hdr[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(txt)
            set_font(r, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            shade(hdr[i], "1F4E79")
        _mark_header_row(t.rows[0])
        _no_split(t.rows[0])

        for row in body:
            cells = t.add_row().cells
            for i in range(ncol):
                val = row[i] if i < len(row) else ""
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.25
                add_rich_text(p, val, size=9.5)
            _no_split(t.rows[-1])

        # 宽度必须同时写到 tblGrid 和每个单元格上,Word 才认
        _set_grid(t, widths)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)

        # 斑马纹
        for ridx, row in enumerate(t.rows[1:], start=1):
            if ridx % 2 == 0:
                for c in row.cells:
                    shade(c, "F2F7FB")

        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0)

    def page_break(self):
        self.doc.add_page_break()


# --------------------------------------------------------------------------
def parse(md_text):
    """切成块序列"""
    lines = md_text.split("\n")
    blocks = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        s = line.rstrip()

        if not s.strip():
            i += 1
            continue

        if s.strip() == "---":
            blocks.append(("break", None))
            i += 1
            continue

        if s.startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            i += 1
            blocks.append(("code", code))
            continue

        if s.startswith("### "):
            blocks.append(("h3", s[4:].strip())); i += 1; continue
        if s.startswith("## "):
            blocks.append(("h2", s[3:].strip())); i += 1; continue
        if s.startswith("# "):
            blocks.append(("h1", s[2:].strip())); i += 1; continue

        if s.lstrip().startswith("|") and s.rstrip().endswith("|"):
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                raw = lines[i].strip()
                cells = [c.strip() for c in raw.strip("|").split("|")]
                rows.append(cells)
                i += 1
            # 去掉分隔行
            cleaned = []
            for r in rows:
                if all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r):
                    continue
                cleaned.append(r)
            if cleaned:
                blocks.append(("table", cleaned))
            continue

        if s.startswith(">"):
            q = []
            while i < n and lines[i].startswith(">"):
                # 裸 ">" 是引用块内的空行,不能当成块结束 ——
                # 否则一段引用会被切成好几个框,中间还漏出一个 ">" 字符。
                content = lines[i][1:]
                if content.startswith(" "):
                    content = content[1:]
                content = content.rstrip()
                if content:
                    q.append(content)
                i += 1
            if q:
                blocks.append(("quote", q))
            continue

        if re.match(r"^- ", s):
            blocks.append(("bullet", s[2:].strip())); i += 1; continue
        if re.match(r"^\d+\. ", s):
            blocks.append(("num", re.sub(r"^\d+\. ", "", s).strip())); i += 1; continue

        blocks.append(("p", s.strip()))
        i += 1

    return blocks


def main():
    src = sys.argv[1]
    out = sys.argv[2]

    with open(src, encoding="utf-8") as f:
        md = f.read()

    blocks = parse(md)

    # 封面:第一个 break 之前的所有块
    cover_blocks = []
    rest = []
    seen_break = False
    for kind, payload in blocks:
        if kind == "break" and not seen_break:
            seen_break = True
            continue
        if not seen_break:
            cover_blocks.append((kind, payload))
        else:
            rest.append((kind, payload))

    b = Builder()
    b.cover(cover_blocks)
    b.page_break()

    for kind, payload in rest:
        if kind == "break":
            b.page_break()
        elif kind == "h1":
            b.h1(payload)
        elif kind == "h2":
            b.h2(payload)
        elif kind == "h3":
            b.h3(payload)
        elif kind == "p":
            b.para(payload)
        elif kind == "bullet":
            b.bullet(payload)
        elif kind == "num":
            b.num(payload)
        elif kind == "code":
            b.code(payload)
        elif kind == "table":
            b.table(payload)
        elif kind == "quote":
            if len(payload) >= 2:
                b.callout(payload[0], "".join(payload[1:]))
            else:
                b.callout(payload[0], "")

    b.doc.save(out)
    print("已生成:", out)
    print("段落数:", len(b.doc.paragraphs), "表格数:", len(b.doc.tables))


if __name__ == "__main__":
    main()
